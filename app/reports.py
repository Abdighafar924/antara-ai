"""
MediCore AI — Report Generators
PDF (ReportLab) and PPTX (python-pptx). No Streamlit dependency — pure functions
that take the cleaned dataframe + computed metrics and return bytes.
"""

import io
import pandas as pd

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                     TableStyle, PageBreak, HRFlowable)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.units import inch
    RL_OK = True
except ImportError:
    RL_OK = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    PPTX_OK = True
except ImportError:
    PPTX_OK = False


def generate_pdf_report(df, cm, quality, hs, maturity, hospital_name="", report_period=""):
    if not RL_OK:
        raise ImportError("reportlab not installed")

    cost_col, los_col, sat_col = cm.get("cost"), cm.get("length_of_stay"), cm.get("satisfaction")
    cond_col, age_col = cm.get("condition"), cm.get("age")
    rr = hs["rr"]

    buf = io.BytesIO()

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(rl_colors.HexColor("#64748b"))
        canvas.drawString(0.9*inch, 0.45*inch, f"{hospital_name} | {report_period} | MediCore AI")
        canvas.drawRightString(letter[0]-0.9*inch, 0.45*inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=0.8*inch, bottomMargin=0.8*inch,
                             leftMargin=0.9*inch, rightMargin=0.9*inch)
    styles = getSampleStyleSheet()
    navy, slate = rl_colors.HexColor("#1a3a6b"), rl_colors.HexColor("#64748b")
    title_style = ParagraphStyle("T", parent=styles["Title"], textColor=navy, fontSize=22, spaceAfter=4)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=navy, fontSize=14, spaceBefore=14, spaceAfter=6)
    body = ParagraphStyle("B", parent=styles["Normal"], fontSize=9.5, leading=14, textColor=rl_colors.HexColor("#334155"))
    small = ParagraphStyle("S", parent=styles["Normal"], fontSize=8.5, leading=13, textColor=rl_colors.HexColor("#334155"))

    def make_table(data, col_widths):
        t = Table(data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),navy), ("TEXTCOLOR",(0,0),(-1,0),rl_colors.white),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"), ("FONTSIZE",(0,0),(-1,-1),8.5),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[rl_colors.HexColor("#f8fafc"), rl_colors.white]),
            ("GRID",(0,0),(-1,-1),0.5,rl_colors.HexColor("#e2e8f0")),
            ("ALIGN",(1,0),(-1,-1),"CENTER"),
            ("TOPPADDING",(0,0),(-1,-1),4), ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ]))
        return t

    story = [Spacer(1,60), Paragraph("MediCore AI", title_style),
             Paragraph("Enterprise Healthcare Intelligence Report", body),
             HRFlowable(width="100%", thickness=2, color=navy, spaceAfter=12), Spacer(1,20)]
    story.append(Paragraph(
        f"<b>{hospital_name or 'Hospital'}</b><br/>Reporting Period: {report_period or 'N/A'}<br/>"
        f"Generated: {pd.Timestamp.now().strftime('%d %B %Y')}<br/>Cohort Size: {len(df):,} patients", body))
    story.append(Spacer(1,20))
    story.append(Paragraph(
        f"Hospital Score: <b>{hs['score']}/100</b> (Grade {hs['grade']}, {hs['risk']} Risk)  |  "
        f"Maturity: <b>{maturity['score']}/100</b> (Grade {maturity['grade']})  |  "
        f"Data Quality: <b>{quality['quality_score']:.0f}/100</b>", body))
    story.append(PageBreak())

    story.append(Paragraph("Executive Summary", h1))
    narrative = [f"This report covers <b>{len(df):,} patients</b>."]
    if age_col in df.columns: narrative.append(f"Average age: {df[age_col].mean():.0f} years.")
    if los_col in df.columns: narrative.append(f"Average LOS: {df[los_col].mean():.1f} days.")
    if "Readmission_Bin" in df.columns: narrative.append(f"Readmission rate: {rr:.1f}% (benchmark: 15%).")
    if sat_col in df.columns: narrative.append(f"Satisfaction: {df[sat_col].mean():.2f}/5 (target: 4.0).")
    if cost_col in df.columns: narrative.append(f"Avg cost: KES {df[cost_col].mean():,.0f}.")
    story.append(Paragraph(" ".join(narrative), body))
    story.append(PageBreak())

    story.append(Paragraph("Key Performance Indicators", h1))
    kpi_data = [["Metric","Value","Benchmark","Status"]]
    if "Readmission_Bin" in df.columns:
        kpi_data.append(["Readmission Rate", f"{rr:.1f}%", "≤15%", "GOOD" if rr<=15 else "ACTION"])
    if sat_col in df.columns:
        sv = df[sat_col].mean()
        kpi_data.append(["Avg Satisfaction", f"{sv:.2f}/5", "≥4.0", "GOOD" if sv>=4 else "WATCH"])
    if cost_col in df.columns:
        cv = df[cost_col].mean()
        kpi_data.append(["Avg Cost", f"KES {cv:,.0f}", "KES 9,000", "GOOD" if cv<=9000 else "ACTION"])
    if los_col in df.columns:
        lv = df[los_col].mean()
        kpi_data.append(["Avg LOS", f"{lv:.1f} d", "≤7 d", "GOOD" if lv<=7 else "WATCH"])
    kpi_data.append(["Data Quality", f"{quality['quality_score']:.0f}/100", "≥85",
                      "GOOD" if quality["quality_score"]>=85 else "WATCH"])
    if len(kpi_data) > 1:
        story.append(make_table(kpi_data, [2.2*inch, 1.3*inch, 1.3*inch, 1.1*inch]))
    story.append(PageBreak())

    story.append(Paragraph("Hospital Performance Score", h1))
    score_data = [["Dimension","Score /100","Weight"]]
    for d, s, w in [("Satisfaction", hs["sat_score"], "30%"), ("Readmission", hs["readm_score"], "25%"),
                    ("LOS", hs["los_score"], "20%"), ("Cost", hs["cost_score"], "15%"),
                    ("Data Quality", hs["quality_score"], "10%"), ("OVERALL", hs["score"], "100%")]:
        score_data.append([d, f"{s:.0f}", w])
    story.append(make_table(score_data, [3*inch, 1.5*inch, 1.2*inch]))
    story.append(PageBreak())

    if cond_col in df.columns:
        story.append(Paragraph("Clinical Summary by Condition", h1))
        agg = {}
        if age_col in df.columns: agg["Patients"] = (age_col, "count")
        if cost_col in df.columns: agg["Avg_Cost"] = (cost_col, "mean")
        if "Readmission_Bin" in df.columns: agg["Readm"] = ("Readmission_Bin", "mean")
        if agg:
            cond_df = df.groupby(cond_col).agg(**agg).reset_index()
            if "Patients" in cond_df: cond_df = cond_df.sort_values("Patients", ascending=False)
            headers = [cond_col] + list(agg.keys())
            tdata = [headers]
            for _, r in cond_df.head(15).iterrows():
                row = [str(r[cond_col])]
                if "Patients" in agg: row.append(str(int(r.get("Patients",0))))
                if "Avg_Cost" in agg: row.append(f"KES {r.get('Avg_Cost',0):,.0f}")
                if "Readm" in agg: row.append(f"{r.get('Readm',0)*100:.1f}%")
                tdata.append(row)
            w = 5.5*inch/len(headers)
            story.append(make_table(tdata, [w]*len(headers)))
        story.append(PageBreak())

    story.append(Paragraph("Strategic Recommendations", h1))
    recs = []
    if "Readmission_Bin" in df.columns and rr > 15:
        recs.append(f"1. READMISSION ({rr:.1f}%): Implement structured post-discharge follow-up. Target <15%.")
    if sat_col in df.columns and df[sat_col].mean() < 4.0:
        recs.append(f"2. SATISFACTION ({df[sat_col].mean():.2f}/5): Hourly rounding, communication training.")
    if cost_col in df.columns and df[cost_col].mean() > 9000:
        recs.append(f"3. COST (KES {df[cost_col].mean():,.0f} avg): Review protocols for LOS optimisation.")
    if quality['quality_score'] < 85:
        recs.append(f"4. DATA QUALITY ({quality['quality_score']:.0f}/100): Address outstanding issues.")
    for r in recs:
        story.append(Paragraph(r, body))
        story.append(Spacer(1,6))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    return buf.read()


def generate_pptx_report(df, cm, quality, hs, maturity):
    if not PPTX_OK:
        raise ImportError("python-pptx not installed")

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

    def add_rect(slide, x, y, w, h, fill, line=None):
        s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
        if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.5)
        else: s.line.fill.background()
        return s

    def add_text(slide, text, x, y, w, h, size=11, bold=False, color="0f172a",
                  align=PP_ALIGN.LEFT, font="Calibri"):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        run = p.add_run(); run.text = str(text)
        run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = rgb(color); run.font.name = font
        return tb

    def slide_header(slide, title, bg="1a3a6b", tc="ffffff"):
        add_rect(slide, 0, 0, 10, 0.75, bg)
        add_text(slide, title, 0.3, 0.1, 9.4, 0.55, size=18, bold=True, color=tc, font="Georgia")

    def kpi_card(slide, x, y, w, h, label, value, sub, accent):
        add_rect(slide, x, y, w, h, "ffffff", "e2e8f0")
        add_rect(slide, x, y, 0.07, h, accent)
        add_text(slide, label, x+0.15, y+0.08, w-0.2, 0.28, size=8, color="64748b")
        add_text(slide, value, x+0.15, y+0.32, w-0.2, 0.65, size=20, bold=True, color="0f172a", font="Georgia")
        add_text(slide, sub, x+0.15, y+h-0.32, w-0.2, 0.28, size=7, color="94a3b8")

    def bar_chart(slide, x, y, w, h, values, labels, colors_list, title=""):
        if not values: return
        max_v = max(values) or 1
        bar_area_h = h - 0.6
        n = len(values)
        bar_w = (w-0.3)/n*0.7
        gap = (w-0.3)/n
        add_text(slide, title, x, y+bar_area_h+0.05, w, 0.28, size=8, bold=True, color="64748b")
        for i,(v,lbl) in enumerate(zip(values, labels)):
            bh = max(0.05, (v/max_v)*bar_area_h)
            bx, by = x+0.15+i*gap, y+bar_area_h-bh
            add_rect(slide, bx, by, bar_w, bh, colors_list[i % len(colors_list)])
            add_text(slide, f"{v:.0f}", bx, by-0.22, bar_w+0.1, 0.22, size=7, color="0f172a", align=PP_ALIGN.CENTER)
            add_text(slide, str(lbl)[:9], bx-0.05, y+bar_area_h+0.05, bar_w+0.1, 0.35, size=6, color="334155", align=PP_ALIGN.CENTER)

    cost_col, los_col, sat_col = cm.get("cost"), cm.get("length_of_stay"), cm.get("satisfaction")
    cond_col = cm.get("condition")
    rr = hs["rr"]
    sc = "0d9488" if hs["score"]>=80 else ("f59e0b" if hs["score"]>=65 else "ef4444")
    pal = ["1a3a6b","0d9488","f59e0b","ef4444","7c3aed","db2777","ea580c","0284c7"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(blank)
    add_rect(s1, 0, 0, 10, 5.625, "07111f")
    add_rect(s1, 0, 0, 0.08, 5.625, "0d9488")
    add_text(s1, "MediCore AI", 0.3, 1.1, 7.5, 0.9, size=34, bold=True, color="f8fafc", font="Georgia")
    add_text(s1, "Enterprise Healthcare Intelligence Suite", 0.3, 2.1, 7.5, 0.5, size=14, color="94a3b8")
    add_text(s1, f"{len(df):,} Patients  |  {pd.Timestamp.now().strftime('%d %B %Y')}", 0.3, 4.9, 9, 0.35, size=8, color="475569")
    add_rect(s1, 7.6, 1.3, 2.1, 2.3, "0f2744", "1a3a6b")
    add_text(s1, str(hs["score"]), 7.6, 1.45, 2.1, 0.9, size=42, bold=True, color=sc, align=PP_ALIGN.CENTER, font="Georgia")
    add_text(s1, "Hospital Score", 7.6, 2.38, 2.1, 0.3, size=8, color="94a3b8", align=PP_ALIGN.CENTER)
    add_text(s1, f"Grade: {hs['grade']}  |  {hs['risk']} Risk", 7.6, 2.72, 2.1, 0.3, size=8, bold=True, color=sc, align=PP_ALIGN.CENTER)

    s2 = prs.slides.add_slide(blank)
    add_rect(s2, 0, 0, 10, 5.625, "f8fafc")
    slide_header(s2, "Executive KPIs")
    kpi_cards = [("Total Patients", f"{len(df):,}", "cohort", "1a3a6b")]
    if "Readmission_Bin" in df.columns:
        kpi_cards.append(("Readmission Rate", f"{rr:.1f}%", "Target ≤15%", "ef4444" if rr>15 else "0d9488"))
    if sat_col and sat_col in df.columns:
        sv = df[sat_col].mean()
        kpi_cards.append(("Avg Satisfaction", f"{sv:.2f}/5", "Target ≥4.0", "0d9488" if sv>=4 else "f59e0b"))
    if cost_col and cost_col in df.columns:
        cv = df[cost_col].mean()
        kpi_cards.append(("Avg Cost", f"KES {cv:,.0f}", "vs KES 9,000", "ef4444" if cv>9000 else "0d9488"))
    if los_col and los_col in df.columns:
        lv = df[los_col].mean()
        kpi_cards.append(("Avg LOS", f"{lv:.1f} d", "Target ≤7 days", "0d9488" if lv<=7 else "f59e0b"))
    kpi_cards.append(("Hospital Score", f"{hs['score']}/100", f"Grade: {hs['grade']}", sc))
    kpi_cards.append(("Maturity", f"{maturity['score']}/100", f"Grade {maturity['grade']}", "7c3aed"))
    for i,(label,val,sub,accent) in enumerate(kpi_cards[:8]):
        col_i, row_i = i%4, i//4
        kpi_card(s2, 0.15+col_i*2.42, 0.9+row_i*2.2, 2.28, 2.0, label, val, sub, accent)

    if cond_col and cond_col in df.columns:
        s3 = prs.slides.add_slide(blank)
        add_rect(s3, 0, 0, 10, 5.625, "ffffff")
        slide_header(s3, "Clinical Analytics")
        cond_agg = df.groupby(cond_col).size().sort_values(ascending=False).head(6)
        bar_chart(s3, 0.2, 0.9, 4.7, 4.0, cond_agg.values.tolist(), cond_agg.index.tolist(), pal, "Patient Volume by Condition")
        if "Readmission_Bin" in df.columns:
            rr_by_cond = (df.groupby(cond_col)["Readmission_Bin"].mean()*100).sort_values(ascending=False).head(6)
            bar_chart(s3, 5.1, 0.9, 4.7, 4.0, rr_by_cond.round(1).values.tolist(), rr_by_cond.index.tolist(),
                      ["ef4444" if v>15 else "f59e0b" for v in rr_by_cond], "Readmission Rate (%) by Condition")

    s4 = prs.slides.add_slide(blank)
    add_rect(s4, 0, 0, 10, 5.625, "f8fafc")
    slide_header(s4, "Strategic Recommendations", "1a3a6b")
    recs_slide = []
    if "Readmission_Bin" in df.columns and rr > 15:
        recs_slide.append(("Readmissions", f"{rr:.1f}% — target <15%. Structured follow-up.", "ef4444"))
    if sat_col and sat_col in df.columns and df[sat_col].mean() < 4.0:
        recs_slide.append(("Satisfaction", f"{df[sat_col].mean():.2f}/5 — hourly rounding, training.", "0d9488"))
    if cost_col and cost_col in df.columns and df[cost_col].mean() > 9000:
        recs_slide.append(("Cost", f"KES {df[cost_col].mean():,.0f} avg — LOS optimisation.", "f59e0b"))
    recs_slide.append(("Data Quality", f"{quality['quality_score']:.0f}/100 — address outliers/missing.", "1a3a6b"))
    for i,(title,b,color) in enumerate(recs_slide[:6]):
        col_i, row_i = i%3, i//3
        x, y = 0.15+col_i*3.28, 0.9+row_i*2.25
        add_rect(s4, x, y, 3.15, 2.1, "ffffff", "e2e8f0")
        add_rect(s4, x, y, 0.07, 2.1, color)
        add_text(s4, title, x+0.15, y+0.1, 3.0, 0.42, size=9, bold=True, color="0f172a")
        add_text(s4, b, x+0.15, y+0.55, 3.0, 1.35, size=8, color="334155")

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
